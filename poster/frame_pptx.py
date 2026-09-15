"""학회 포스터 프레임 (95 × 120 cm, 세로) — 글씨·그림을 넣을 빈 틀만 만든다.

한국에너지기후변화학회 참고 양식(`data/…포스터 V5.pdf`)의 배치를 따른다: 제목 띠 → 2단 본문
(왼쪽: 초록 · 연구 방법 · 연구 내용 / 오른쪽: 설비·열전달식 패널 · 연구 결과 · 결론 · 사사 · 참고문헌).
모든 글상자는 PowerPoint 에서 바로 고쳐 쓸 수 있고, 회색 안내 문구를 지우고 내용을 넣으면 된다.

    .venv/bin/python poster/frame_pptx.py      # → ttttt/포스터_프레임_95x120cm.pptx + 미리보기 png

같은 배치 명세로 PPTX 와 미리보기(matplotlib)를 함께 그려, LibreOffice 없이 모양을 확인한다.
"""
from __future__ import annotations

import pathlib

import matplotlib.patches as mpatches
import matplotlib.pyplot as plt
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

ROOT = pathlib.Path(__file__).resolve().parents[1]
OUT = ROOT / "ttttt"
FONT = "맑은 고딕"                  # PowerPoint(Windows) 기본 한글 글꼴
W, H = 95.0, 120.0
M, G = 2.0, 1.5                    # 바깥 여백 · 단 사이 간격 [cm]
CW = (W - 2 * M - G) / 2           # 단 너비 44.75 cm
XL, XR = M, M + CW + G

# 색 — 남색 제목 띠, 파랑 섹션 머리, 옅은 패널. 방법 박스는 파스텔 3색, 수식 띠는 보라.
NAVY, BLUE, BLUE_D = "1D3557", "2A78D6", "1C5CAB"
PANEL, PANEL_LINE = "F5F8FC", "C9D6E8"
INK, HINT, WHITE = "0B0B0B", "7A8594", "FFFFFF"
PINK, YELLOW, SKY = "FDE8EC", "FFF4CC", "E3EFFF"
PURPLE, PURPLE_L = "5B4B9E", "EEEAF8"
TEAL_L = "E3F4F1"
FIG, FIG_LINE, CAPTION = "EEF1F5", "9AA5B1", "FFF3A3"

items: list[dict] = []


def box(x, y, w, h, fill=None, line=None, lw=0.0, dash=False, shape="rect", radius=0.0, text="", size=28,
        bold=False, color=INK, align="left", anchor="top", pad=0.6):
    items.append(dict(x=x, y=y, w=w, h=h, fill=fill, line=line, lw=lw, dash=dash, shape=shape, radius=radius,
                      text=text, size=size, bold=bold, color=color, align=align, anchor=anchor, pad=pad))


def section(x, y, w, title, panel_h, head_h=3.0, size=54):
    """섹션 머리 띠 + 아래 패널. 패널 안쪽 시작 y 를 돌려준다."""
    box(x, y, w, head_h, fill=BLUE, shape="round", radius=0.35, text=title, size=size, bold=True, color=WHITE,
        align="center", anchor="middle", pad=0.2)
    box(x, y + head_h + 0.3, w, panel_h, fill=PANEL, line=PANEL_LINE, lw=0.08, shape="round", radius=0.02)
    return y + head_h + 0.3


def figure(x, y, w, h, label="그림 영역", cap="그림 캡션 (짧은 명사구)", cap_h=1.7):
    box(x, y, w, h, fill=FIG, line=FIG_LINE, lw=0.06, dash=True, text=f"{label}\n(삽입 → 그림)", size=26,
        color=HINT, align="center", anchor="middle")
    if cap:
        box(x + w * 0.15, y + h + 0.25, w * 0.7, cap_h, fill=CAPTION, shape="round", radius=0.25, text=cap, size=26,
            bold=True, color=INK, align="center", anchor="middle", pad=0.2)


def text(x, y, w, h, hint, size=30):
    box(x, y, w, h, text=hint, size=size, color=HINT)


# ───────────────────────── 제목 띠
box(0, 0, W, 15.5, fill=NAVY)
box(0, 15.5, W, 0.5, fill=BLUE)
box(M, 1.4, 66, 5.4, text="국문 제목을 입력하세요 (1~2줄)", size=88, bold=True, color=WHITE, anchor="middle")
box(M, 6.9, 66, 2.3, text="English title of the poster", size=44, color="DCE6F2", anchor="middle")
box(M, 9.5, 66, 2.1, text="저자1, 저자2, 저자3, 교신저자*", size=40, bold=True, color=WHITE, anchor="middle")
box(M, 11.7, 66, 2.1, text="소속, 주소, 우편번호", size=32, color="DCE6F2", anchor="middle")
for i, lab in enumerate(("학회 로고", "기관 로고")):
    box(70.2 + i * 12.4, 2.8, 10.8, 9.8, fill=WHITE, line="B8C4D6", lw=0.06, dash=True, shape="round", radius=0.08,
        text=lab, size=30, color=HINT, align="center", anchor="middle")

# ───────────────────────── 왼쪽 단
# 1) 초록
y = section(XL, 17.5, CW, "초 록 (Abstract)", 11.5)
text(XL + 1, y + 0.8, CW - 2, 10, "초록 내용을 입력하세요 (8~10줄 · 연구 배경 → 방법 → 결과 → 의의)")

# 2) 연구 방법
y = section(XL, 33.8, CW, "연구 방법 (Methodology)", 38.0)
bw = (CW - 2 - 2 * 1.0) / 3
for i, (fill, t) in enumerate(((PINK, "1. 방법 제목  기호"), (YELLOW, "2. 방법 제목  기호"), (SKY, "3. 방법 제목  기호"))):
    bx = XL + 1 + i * (bw + 1.0)
    box(bx, y + 1.0, bw, 12.0, fill=fill, shape="round", radius=0.06)
    box(bx, y + 1.3, bw, 2.6, text=t, size=34, bold=True, color=INK, anchor="middle")
    text(bx, y + 4.2, bw, 8.4, "- 핵심 내용 1\n- 핵심 내용 2\n- 핵심 내용 3", size=26)
figure(XL + 1, y + 14.0, CW - 2, 8.5, label="흐름도 영역 (도형 · 화살표)", cap=None)
box(XL + 1, y + 23.5, CW - 2, 2.4, fill=PURPLE, shape="round", radius=0.3, text="핵심 수식", size=36, bold=True,
    color=WHITE, align="center", anchor="middle", pad=0.1)
box(XL + 1, y + 26.1, CW - 2, 7.0, fill=WHITE, line=PURPLE, lw=0.08, shape="round", radius=0.05,
    text="수식을 입력하세요 (삽입 → 수식)", size=34, color=HINT, align="center", anchor="middle")
box(XL + 1, y + 33.5, CW - 2, 3.6, fill=PURPLE_L, shape="round", radius=0.1,
    text="기호 정의:  A = …,  B = …,  C = …", size=26, color=HINT, anchor="middle")

# 3) 연구 내용
y = section(XL, 76.6, CW, "연구 내용 (Experimental details)", 38.1)
text(XL + 1, y + 1.0, 21.4, 11.5, "- 굵은 머리말: 설명 한 문장\n- 굵은 머리말: 설명 한 문장\n- 굵은 머리말: 설명 한 문장\n- 굵은 머리말: 설명 한 문장", size=28)
figure(XL + 23.4, y + 1.0, CW - 24.4, 9.6)
fw = (CW - 2 - 1.2) / 2
figure(XL + 1, y + 14.2, fw, 9.8)
figure(XL + 2.2 + fw, y + 14.2, fw, 9.8)
box(XL + 1, y + 27.6, CW - 2, 9.8, fill=FIG, line=FIG_LINE, lw=0.06, dash=True,
    text="표 영역 (삽입 → 표)", size=28, color=HINT, align="center", anchor="middle")

# ───────────────────────── 오른쪽 단
# 4) 설비 구조 · 열전달식 (계산 패널)
y = section(XR, 17.5, CW, "설비 구조 · 열전달식", 27.0)
figure(XR + 1, y + 1.0, 26.0, 12.2, label="설비 구조도 영역", cap=None)
box(XR + 28.0, y + 1.0, CW - 29.0, 12.2, fill=FIG, line=FIG_LINE, lw=0.06, dash=True,
    text="판정 표 영역\n(식 · 성립 여부)", size=26, color=HINT, align="center", anchor="middle")
ew = (CW - 2 - 2 * 2.4) / 3
for i in range(3):
    ex = XR + 1 + i * (ew + 2.4)
    box(ex, y + 14.4, ew, 6.8, fill=WHITE, line=BLUE, lw=0.08, shape="round", radius=0.08)
    box(ex, y + 14.6, ew, 2.0, text=f"식 {'①②③'[i]} 이름", size=30, bold=True, color=BLUE_D, align="center", anchor="middle")
    box(ex, y + 16.7, ew, 4.3, text="수식 · 값", size=26, color=HINT, align="center", anchor="middle")
    if i < 2:
        box(ex + ew + 0.4, y + 16.6, 1.6, 2.4, fill=BLUE, shape="arrow")
box(XR + 1, y + 22.3, CW - 2, 3.8, fill=TEAL_L, shape="round", radius=0.1,
    text="요점: 한 문장으로 이 패널의 결론을 입력하세요", size=28, bold=True, color=HINT, anchor="middle")

# 5) 연구 결과
y = section(XR, 48.3, CW, "연구 결과 (Research Results)", 36.8)
text(XR + 1, y + 1.0, 20.9, 10.5, "- 요점 1: 숫자를 포함한 한 문장\n- 요점 2: 숫자를 포함한 한 문장\n- 요점 3: 숫자를 포함한 한 문장", size=28)
figure(XR + 22.85, y + 1.0, 20.9, 8.6)
figure(XR + 1, y + 12.6, CW - 2, 10.8, label="시계열 그림 영역 (가로로 넓게)")
figure(XR + 1, y + 26.6, 20.9, 8.0)
box(XR + 22.85, y + 26.6, 20.9, 9.8, fill=FIG, line=FIG_LINE, lw=0.06, dash=True,
    text="표 영역 (삽입 → 표)", size=28, color=HINT, align="center", anchor="middle")

# 6) 결론
y = section(XR, 89.9, CW, "결 론 (Conclusion)", 10.8)
text(XR + 1, y + 0.8, CW - 2, 9.4, "· 결론 1 (숫자 포함)\n· 결론 2\n· 결론 3\n· 결론 4\n· 결론 5", size=30)

# 7) 사사 · 8) 참고문헌 — 작은 머리 띠
y = section(XR, 104.7, CW, "사 사 (Acknowledgement)", 3.1, head_h=2.4, size=40)
text(XR + 1, y + 0.2, CW - 2, 2.8, "과제명 · 과제번호를 입력하세요", size=24)
y = section(XR, 111.6, CW, "참고문헌 (Reference)", 3.7, head_h=2.4, size=40)
text(XR + 1, y + 0.1, CW - 2, 3.5, "[1] 저자, \"제목\", 학술지, 권(호), 쪽, 연도.\n[2] …", size=22)


# ───────────────────────── PPTX
def _rgb(h):
    return RGBColor.from_string(h)


def build_pptx(path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(W), Cm(H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    kinds = {"rect": MSO_SHAPE.RECTANGLE, "round": MSO_SHAPE.ROUNDED_RECTANGLE, "arrow": MSO_SHAPE.RIGHT_ARROW}
    for it in items:
        needs_shape = it["fill"] or it["line"] or it["shape"] != "rect"
        if needs_shape:
            sh = slide.shapes.add_shape(kinds[it["shape"]], Cm(it["x"]), Cm(it["y"]), Cm(it["w"]), Cm(it["h"]))
            if it["shape"] == "round":
                sh.adjustments[0] = min(max(it["radius"], 0.01), 0.5)
            if it["fill"]:
                sh.fill.solid(); sh.fill.fore_color.rgb = _rgb(it["fill"])
            else:
                sh.fill.background()
            if it["line"]:
                sh.line.color.rgb = _rgb(it["line"]); sh.line.width = Cm(it["lw"])
                if it["dash"]:
                    sh.line.dash_style = MSO_LINE.DASH
            else:
                sh.line.fill.background()
            sh.shadow.inherit = False
        else:
            sh = slide.shapes.add_textbox(Cm(it["x"]), Cm(it["y"]), Cm(it["w"]), Cm(it["h"]))
        if not it["text"]:
            continue
        tf = sh.text_frame
        tf.word_wrap = True
        tf.auto_size = None
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, side, Cm(it["pad"]))
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE}[it["anchor"]]
        for i, line in enumerate(it["text"].split("\n")):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER}[it["align"]]
            run = para.add_run()
            run.text = line
            f = run.font
            f.size, f.bold, f.name = Pt(it["size"]), it["bold"], FONT
            f.color.rgb = _rgb(it["color"])
            rpr = run._r.get_or_add_rPr()                       # 한글(동아시아) 글꼴도 지정
            latin = rpr.find(qn("a:latin"))
            ea = etree.SubElement(rpr, qn("a:ea")); ea.set("typeface", FONT)
            if latin is not None:
                latin.addnext(ea)
    prs.save(path)


# ───────────────────────── 미리보기 (같은 명세)
def build_preview(path):
    plt.rcParams["font.family"] = "Noto Sans CJK KR"
    fig = plt.figure(figsize=(W / 2.54, H / 2.54), dpi=30)
    ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, W); ax.set_ylim(H, 0); ax.axis("off")
    ax.add_patch(mpatches.Rectangle((0, 0), W, H, color="#FFFFFF"))
    for it in items:
        x, y, w, h = it["x"], it["y"], it["w"], it["h"]
        kw = dict(facecolor=f"#{it['fill']}" if it["fill"] else "none",
                  edgecolor=f"#{it['line']}" if it["line"] else "none",
                  linewidth=it["lw"] / 2.54 * 72 if it["line"] else 0, linestyle="--" if it["dash"] else "-")
        if it["shape"] == "round":
            r = it["radius"] * min(w, h)
            ax.add_patch(mpatches.FancyBboxPatch((x, y), w, h, boxstyle=f"round,pad=0,rounding_size={r}", **kw))
        elif it["shape"] == "arrow":
            ax.add_patch(mpatches.FancyArrow(x, y + h / 2, w, 0, width=h * 0.5, head_width=h, head_length=w * 0.45,
                                             length_includes_head=True, color=f"#{it['fill']}"))
        elif it["fill"] or it["line"]:
            ax.add_patch(mpatches.Rectangle((x, y), w, h, **kw))
        if it["text"]:
            tx = x + w / 2 if it["align"] == "center" else x + it["pad"]
            ty = y + h / 2 if it["anchor"] == "middle" else y + it["pad"]
            ax.text(tx, ty, it["text"], fontsize=it["size"], fontweight="bold" if it["bold"] else "normal",
                    color=f"#{it['color']}", ha="center" if it["align"] == "center" else "left",
                    va="center" if it["anchor"] == "middle" else "top", linespacing=1.25)
    fig.savefig(path, dpi=30)
    plt.close(fig)


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    build_pptx(OUT / "포스터_프레임_95x120cm.pptx")
    build_preview(OUT / "포스터_프레임_미리보기.png")
    print("저장:", OUT / "포스터_프레임_95x120cm.pptx", "·", OUT / "포스터_프레임_미리보기.png", "· 도형", len(items))
